"""
services/extraction_service.py
──────────────────────────────
Logic for converting raw files (PDF/Images/CSV/XLSX/XML) into structured Markdown 
and JSON using Azure Document Intelligence or native parsers.
"""
from typing import Dict, Any, Tuple
import os
import base64
from datetime import datetime, timezone
from io import StringIO, BytesIO

import csv
import openpyxl
import xml.etree.ElementTree as ET
from sqlalchemy.orm import Session
from azure.ai.documentintelligence import DocumentIntelligenceClient
from azure.ai.documentintelligence.models import AnalyzeDocumentRequest
from azure.core.credentials import AzureKeyCredential

from core.config import settings
from core.logger import get_logger
from services.blob_service import BlobService
from models.document import Document, DocumentPage
from models.source_mapping import SourceMapping

logger = get_logger(__name__)

def _get_di_client() -> DocumentIntelligenceClient:
    """Initializes the Azure Document Intelligence client."""
    if not settings.AZURE_DI_ENDPOINT or not settings.AZURE_DI_KEY:
        raise ValueError("Azure Document Intelligence credentials not configured.")
    
    return DocumentIntelligenceClient(
        endpoint=settings.AZURE_DI_ENDPOINT,
        credential=AzureKeyCredential(settings.AZURE_DI_KEY)
    )

def _parse_csv(db: Session, document_id: str, file_bytes: bytes) -> str:
    """Native CSV parser to Markdown table."""
    csv_text = file_bytes.decode("utf-8", errors="replace")
    reader = csv.reader(StringIO(csv_text))
    rows = list(reader)
    if not rows: return "Empty CSV"
    
    mappings = []
    md = ["| " + " | ".join(f"{c} [[r0_c{i}]]" for i, c in enumerate(rows[0])) + " |", "|" + "|".join(["---"] * len(rows[0])) + "|"]
    for i, c in enumerate(rows[0]):
        mappings.append(SourceMapping(document_id=document_id, ref_key=f"r0_c{i}", file_type="csv", row_index=0, col_index=i))
    
    for r_idx, row in enumerate(rows[1:], start=1):
        md.append("| " + " | ".join(f"{c} [[r{r_idx}_c{i}]]" for i, c in enumerate(row)) + " |")
        for i, c in enumerate(row):
            mappings.append(SourceMapping(document_id=document_id, ref_key=f"r{r_idx}_c{i}", file_type="csv", row_index=r_idx, col_index=i))
    
    db.add_all(mappings)
    return "\n".join(md)

def _parse_xlsx(db: Session, document_id: str, file_bytes: bytes) -> Tuple[str, str]:
    """Native XLSX parser to Markdown table."""
    wb = openpyxl.load_workbook(BytesIO(file_bytes), data_only=True)
    sheet = wb.active
    rows = []
    for row in sheet.iter_rows(values_only=True):
        clean_row = [str(cell) if cell is not None else "" for cell in row]
        if any(clean_row): rows.append(clean_row)
    if not rows: return "Empty Sheet", sheet.title
    
    mappings = []
    md = ["| " + " | ".join(f"{c} [[r0_c{i}]]" for i, c in enumerate(rows[0])) + " |", "|" + "|".join(["---"] * len(rows[0])) + "|"]
    for i, c in enumerate(rows[0]):
        mappings.append(SourceMapping(document_id=document_id, ref_key=f"r0_c{i}", file_type="xlsx", row_index=0, col_index=i))
        
    for r_idx, row in enumerate(rows[1:], start=1):
        md.append("| " + " | ".join(f"{c} [[r{r_idx}_c{i}]]" for i, c in enumerate(row)) + " |")
        for i, c in enumerate(row):
            mappings.append(SourceMapping(document_id=document_id, ref_key=f"r{r_idx}_c{i}", file_type="xlsx", row_index=r_idx, col_index=i))
            
    db.add_all(mappings)
    return "\n".join(md), sheet.title

def _parse_xml(db: Session, document_id: str, file_bytes: bytes) -> str:
    """Native XML parser to Markdown table (best effort for flat structures)."""
    xml_text = file_bytes.decode("utf-8", errors="replace")
    rows = []
    try:
        root = ET.fromstring(xml_text)
        if len(root) > 0:
            headers = [child.tag for child in root[0]]
            if headers:
                rows.append(headers)
                for item in root:
                    rows.append([(item.find(h).text or "").strip() if item.find(h) is not None else "" for h in headers])
    except Exception: pass
    if not rows: return f"```xml\n{xml_text}\n```"
    
    mappings = []
    md = ["| " + " | ".join(f"{c} [[xpath_0_c{i}]]" for i, c in enumerate(rows[0])) + " |", "|" + "|".join(["---"] * len(rows[0])) + "|"]
    for i, c in enumerate(rows[0]):
        mappings.append(SourceMapping(document_id=document_id, ref_key=f"xpath_0_c{i}", file_type="xml", row_index=0, col_index=i))
        
    for r_idx, row in enumerate(rows[1:], start=1):
        md.append("| " + " | ".join(f"{c} [[xpath_{r_idx}_c{i}]]" for i, c in enumerate(row)) + " |")
        for i, c in enumerate(row):
            mappings.append(SourceMapping(document_id=document_id, ref_key=f"xpath_{r_idx}_c{i}", file_type="xml", row_index=r_idx, col_index=i))

    db.add_all(mappings)
    return "\n".join(md)

def _enrich_azure_markdown(db: Session, document_id: str, result: Any, original_content: str) -> str:
    """
    Injects [[ref_N]] tags into Azure's Markdown output and saves mappings.
    Uses 'paragraphs' from Azure result to find text spans.
    """
    if not result.paragraphs:
        return original_content

    # Sort paragraphs by their span start to avoid offset shifts
    # However, we'll work backwards or use a list/joining approach
    enriched_content = original_content
    
    # We want to insert tags at the end of each paragraph or line
    # For now, we'll use paragraphs as the most reliable anchor
    
    # Store mappings in a list to bulk insert later
    mappings = []
    
    # Work from the end of the content to the beginning to keep spans valid
    # But Azure DI model-layout paragraphs spans are 0-based index in 'original_content'
    sorted_paragraphs = sorted(result.paragraphs, key=lambda p: p.spans[0].offset, reverse=True)
    
    for i, para in enumerate(sorted_paragraphs):
        # We need an index from the original list (or just a counter)
        ref_key = f"ref_{len(result.paragraphs) - 1 - i}"
        
        # Capture physical location from the first span (usually paragraphs are single span)
        page_num = para.bounding_regions[0].page_number if para.bounding_regions else 1
        # Bounding polygon is [x1, y1, x2, y2, x3, y3, x4, y4]
        bbox = para.bounding_regions[0].polygon if para.bounding_regions else None
        
        mappings.append(SourceMapping(
            document_id=document_id,
            ref_key=ref_key,
            file_type="pdf",
            page_number=page_num,
            bbox=bbox
        ))
        
        # Inject the tag: find the end of the span and insert it
        # Actually, it's safer to insert it right AFTER the text in the markdown
        offset = para.spans[0].offset + para.spans[0].length
        enriched_content = enriched_content[:offset] + f" [[{ref_key}]]" + enriched_content[offset:]

    db.add_all(mappings)
    db.commit()
    return enriched_content

def process_document(db: Session, document_id: str) -> Dict[str, Any]:
    """
    Orchestrates the extraction pipeline:
    Automatically selects between Azure DI and native parsers based on file extension.
    """
    logger.info("Starting document extraction", extra={"doc_id": document_id})
    
    doc = db.query(Document).filter(Document.id == document_id).first()
    if not doc:
        raise ValueError(f"Document {document_id} not found.")

    doc.status = "di_processing"
    db.commit()

    blob_service = BlobService()
    di_client = _get_di_client()

    try:
        # 1. Download from Blob
        file_bytes = blob_service.download_bytes(doc.blob_path)
        extension = os.path.splitext(doc.filename)[1].lower() if doc.filename else ""
        
        markdown_content = ""
        operation_id = None
        page_count = 0
        table_count = 0

        # 2. Extract content
        if extension == ".csv":
            markdown_content = _parse_csv(db, document_id, file_bytes)
            page_count, table_count = 1, 1
        elif extension == ".xlsx":
            md, title = _parse_xlsx(db, document_id, file_bytes)
            markdown_content = f"## Sheet: {title}\n\n" + md
            page_count, table_count = 1, 1
        elif extension == ".xml":
            markdown_content = _parse_xml(db, document_id, file_bytes)
            page_count, table_count = 1, 1
        else:
            # Azure DI for PDF, DOCX, Images
            # SDK requires bytes_source to be base64 encoded for certain models or versions
            b64_content = base64.b64encode(file_bytes).decode("utf-8")
            request = AnalyzeDocumentRequest(bytes_source=b64_content)
            
            poller = di_client.begin_analyze_document(
                model_id="prebuilt-layout",
                body=request,
                output_content_format="markdown",
            )
            result = poller.result()
            
            # Traceability Enrichment for PDF/Images
            markdown_content = _enrich_azure_markdown(db, document_id, result, result.content)
            
            operation_id = getattr(result, "operation_id", None)
            page_count = len(result.pages) if result.pages else 0
            table_count = len(result.tables) if result.tables else 0

        # 3. Save Markdown to Blob
        processed_path = blob_service.processed_path(document_id)
        blob_service.upload_text(
            text=markdown_content,
            blob_path=processed_path,
            content_type="text/markdown"
        )

        # 4. Update Database records
        doc.status = "di_processed"
        doc.processed_at = datetime.now(timezone.utc)
        
        db.query(DocumentPage).filter(DocumentPage.document_id == document_id).delete()

        for i in range(1, page_count + 1):
            db_page = DocumentPage(
                document_id=document_id,
                page_number=i,
                page_blob_path=processed_path,
                azure_operation_id=operation_id
            )
            db.add(db_page)

        doc.stats = doc.stats or {}
        doc.stats["page_count"] = page_count
        doc.stats["table_count"] = table_count
        
        db.commit()
        logger.info("Document extraction complete", extra={"doc_id": document_id})
        
        return {
            "status": "success",
            "document_id": document_id,
            "pages": page_count
        }

    except Exception as e:
        logger.error(f"Extraction failed: {str(e)}", extra={"doc_id": document_id})
        doc.status = "failed"
        db.commit()
        raise e
